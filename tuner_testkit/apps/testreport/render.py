"""Report Markdown (assets/testreport/...) -> GJB438B docx through the docxtpl template."""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from . import mdparse
from .draft import TODO_MARK

REVISION_COLUMNS = {"序号": "seq", "版本号": "version", "修改日期": "date", "修改内容": "change", "修改人": "author"}
REFERENCE_COLUMNS = {"序号": "seq", "文档编号": "code", "文档标题": "title", "编写单位": "org", "发布日期": "date"}
RESIDUAL_COLUMNS = {"测试用例标识": "case", "描述": "desc"}
SOFTWARE_COLUMNS = {"序号": "seq", "支持软件名称": "name", "版本号": "version", "用途": "purpose", "介质": "media", "来源": "source", "保密性": "secrecy"}
HARDWARE_COLUMNS = {"序号": "seq", "硬件名称": "name", "配置": "config", "推荐型号": "model", "用途": "purpose", "数量": "qty", "保密性": "secrecy"}
IMPACT_COLUMNS = {"序号": "seq", "测试环境": "test_env", "实装环境": "real_env", "影响分析": "analysis", "解决措施": "measure"}
EXEC_COLUMNS = {"测试类型": "type", "设计用例": "designed", "执行/未执行": "run", "通过/未通过": "passed"}
DEFECT_COLUMNS = {
    "严重性等级": "level", "需求缺陷": "req", "设计缺陷": "design", "编码缺陷": "code", "文档缺陷": "doc",
    "其它缺陷": "other", "归零缺陷数": "closed", "遗留缺陷数": "open",
}
CASE_COLUMNS = {
    "序号": "seq", "模块": "module", "子模块": "suite", "用例编号（测试标识）": "id", "测试用例名称": "title",
    "前置条件": "pre", "测试步骤": "action", "预期结果": "expected", "结果": "result",
}
SUMMARY_KEYS = {
    "软件标识号": "software_id", "软件名称": "software_name", "编程语言": "language", "程序规模": "size",
    "测试周期": "period", "测试类型": "test_type", "测试人员": "testers", "备注": "remark",
}
COUNT_KEYS = ("designed", "executed", "passed", "failed", "not_run")


class ReportFormatError(ValueError):
    pass


def _section(parts: dict[str, list[mdparse.Block]], title: str) -> list[mdparse.Block]:
    if title not in parts:
        raise ReportFormatError(f"报告缺少章节「{title}」")
    return parts[title]


def _tables(blocks: list[mdparse.Block]) -> list[mdparse.Table]:
    return [b.table for b in blocks if b.kind == "table" and b.table is not None]


def _kv(blocks: list[mdparse.Block], title: str) -> dict[str, str]:
    for table in _tables(blocks):
        if table.header[:2] == ["项", "内容"]:
            return {row[0]: mdparse.plain(row[1]) if len(row) > 1 else "" for row in table.rows if row}
    raise ReportFormatError(f"「{title}」缺少「项 | 内容」表")


def _records(table: mdparse.Table, columns: dict[str, str], title: str) -> list[dict[str, str]]:
    missing = [name for name in columns if name not in table.header]
    if missing:
        raise ReportFormatError(f"「{title}」表缺少列：{'、'.join(missing)}")
    return [{key: mdparse.plain(record.get(name, "")) for name, key in columns.items()} for record in table.records()]


def _first_table(blocks: list[mdparse.Block], title: str, first_header: str | None = None) -> mdparse.Table:
    for table in _tables(blocks):
        if first_header is None or (table.header and table.header[0] == first_header):
            return table
    raise ReportFormatError(f"「{title}」缺少表格")


def _paras(blocks: list[mdparse.Block]) -> list[str]:
    return [mdparse.plain(b.text) for b in blocks if b.kind == "para" and b.text.strip()]


def _split_around_table(blocks: list[mdparse.Block]) -> tuple[list[str], list[str]]:
    before: list[str] = []
    after: list[str] = []
    seen = False
    for block in blocks:
        if block.kind == "table":
            seen = True
        elif block.text.strip():
            (after if seen else before).append(mdparse.plain(block.text))
    return before, after


def _need(mapping: dict[str, str], key: str, title: str) -> str:
    if key not in mapping:
        raise ReportFormatError(f"「{title}」缺少「{key}」")
    return mapping[key]


def parse_report(path: Path) -> dict[str, Any]:
    text = path.read_text(encoding="utf-8")
    meta, body = mdparse.split_front_matter(text)
    todos = body.count(TODO_MARK)
    if todos:
        raise ReportFormatError(f"报告还有 {todos} 处「{TODO_MARK} … -->」没有处理")
    parts = mdparse.outline(body)

    cover = _kv(_section(parts, "封面"), "封面")
    ident = _kv(_section(parts, "标识"), "标识")
    parties = _kv(_section(parts, "供需方情况"), "供需方情况")
    evaluation = _kv(_section(parts, "对被测试软件的总体评估"), "对被测试软件的总体评估")
    summary_blocks = _section(parts, "测试结果总结")
    summary = _kv(summary_blocks, "测试结果总结")
    counts = _first_table(summary_blocks, "测试结果总结", "轮次")
    count_rows = {row[0]: row[1:] for row in counts.rows if row}
    exec_blocks = _section(parts, "测试执行情况")
    execution, regression = _split_around_table(exec_blocks)
    defect_blocks = _section(parts, "缺陷归零情况")
    defect_rows = _records(_first_table(defect_blocks, "缺陷归零情况", "严重性等级"), DEFECT_COLUMNS, "缺陷归零情况")

    def counts_of(label: str) -> dict[str, str]:
        if label not in count_rows:
            raise ReportFormatError(f"「测试结果总结」轮次表缺少「{label}」行")
        values = count_rows[label]
        return {key: mdparse.plain(values[idx]) if idx < len(values) else "" for idx, key in enumerate(COUNT_KEYS)}

    first_round = counts_of("测试用例")
    regression_counts = counts_of("回归测试用例")
    regression_counts["reuse"] = regression_counts.pop("designed")

    totals = next((row for row in defect_rows if row["level"] == "合计"), None)
    if totals is None:
        raise ReportFormatError("「缺陷归零情况」表缺少「合计」行")

    return {
        "meta": meta,
        "system_name": _need(cover, "系统名称", "封面"),
        "company": _need(cover, "编制单位", "封面"),
        "report_month": _need(cover, "编制日期", "封面"),
        "revisions": _records(_first_table(_section(parts, "修订页"), "修订页"), REVISION_COLUMNS, "修订页"),
        "doc_title": _need(ident, "文档标题", "标识"),
        "doc_no": ident.get("文档标识号", ""),
        "doc_version": _need(ident, "文档版本号", "标识"),
        "abbreviations": ident.get("缩略语", "无"),
        "overview": _paras(_section(parts, "系统概述")),
        "buyer": parties.get("需方", ""),
        "developer": parties.get("开发方", ""),
        "support": parties.get("保障机构", ""),
        "run_site": _paras(_section(parts, "运行现场")),
        "doc_overview": _paras(_section(parts, "文档概述")),
        "references": _records(_first_table(_section(parts, "引用文档"), "引用文档"), REFERENCE_COLUMNS, "引用文档"),
        "version": _need(evaluation, "软件版本", "对被测试软件的总体评估"),
        "verdict": _need(evaluation, "结论", "对被测试软件的总体评估"),
        "verdict_detail": evaluation.get("结论说明", ""),
        "residual_intro": evaluation.get("残留问题说明", ""),
        "residuals": _records(_first_table(_section(parts, "残留问题"), "残留问题"), RESIDUAL_COLUMNS, "残留问题"),
        "evaluation": _paras(_section(parts, "评估说明")),
        "software": _records(_first_table(_section(parts, "软件环境"), "软件环境"), SOFTWARE_COLUMNS, "软件环境"),
        "hardware": _records(_first_table(_section(parts, "硬件环境"), "硬件环境"), HARDWARE_COLUMNS, "硬件环境"),
        "env_impacts": _records(_first_table(_section(parts, "测试环境影响"), "测试环境影响"), IMPACT_COLUMNS, "测试环境影响"),
        "suggestions": _paras(_section(parts, "改进建议")),
        "s": {key: summary.get(label, "") for label, key in SUMMARY_KEYS.items()},
        "fr": first_round,
        "rg": regression_counts,
        "execution": execution,
        "exec_rows": _records(_first_table(exec_blocks, "测试执行情况", "测试类型"), EXEC_COLUMNS, "测试执行情况"),
        "regression": regression,
        "defect_summary": _paras(defect_blocks),
        "defect_rows": [row for row in defect_rows if row["level"] != "合计"],
        "dt": totals,
        "notes": _paras(_section(parts, "注释")),
        "cases": _records(_first_table(_section(parts, "测试用例"), "测试用例"), CASE_COLUMNS, "测试用例"),
    }


def _joined(paragraphs: list[str], empty: str = "") -> str:
    return "\a".join(paragraphs) if paragraphs else empty


def to_context(parsed: dict[str, Any]) -> dict[str, Any]:
    """Join paragraph lists with docxtpl's paragraph break and wrap every string in Listing."""

    from docxtpl import Listing

    flat = dict(parsed)
    flat.pop("meta", None)
    for key, empty in (
        ("overview", ""), ("run_site", ""), ("doc_overview", ""), ("evaluation", ""), ("suggestions", "无。"),
        ("execution", ""), ("regression", ""), ("defect_summary", ""), ("notes", "无。"),
    ):
        flat[key] = _joined(flat[key], empty)
    if not flat["residuals"]:
        flat["residuals"] = [{"case": "/", "desc": "/"}]

    def wrap(value: Any) -> Any:
        if isinstance(value, str):
            return Listing(value)
        if isinstance(value, dict):
            return {k: wrap(v) for k, v in value.items()}
        if isinstance(value, list):
            return [wrap(v) for v in value]
        return value

    return wrap(flat)


_TAG_RE = re.compile(r"\{\{|\}\}|\{%|%\}")


def check_docx(path: Path) -> dict[str, int]:
    """Re-open the rendered docx: count leftover template tags and rows in the case appendix."""

    from docx import Document

    doc = Document(str(path))
    texts = [p.text for p in doc.paragraphs]
    case_rows = -1
    for table in doc.tables:
        header = [cell.text for cell in table.rows[0].cells] if table.rows else []
        if "用例编号（测试标识）" in header:
            case_rows = len(table.rows) - 1
        texts += [cell.text for row in table.rows for cell in row.cells]
    for section in doc.sections:
        for part in (section.header, section.footer):
            texts += [p.text for p in part.paragraphs]
    return {"leftover_tags": sum(len(_TAG_RE.findall(t)) for t in texts), "case_rows": case_rows}


def render_docx(report: Path, template: Path, out: Path) -> dict[str, Any]:
    from docxtpl import DocxTemplate

    parsed = parse_report(report)
    doc = DocxTemplate(str(template))
    doc.render(to_context(parsed), autoescape=True)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    return {
        "cases": len(parsed["cases"]),
        "open_defects": parsed["dt"].get("open", ""),
        "verdict": parsed["verdict"],
        "id": parsed["meta"].get("id", ""),
        **check_docx(out),
    }
