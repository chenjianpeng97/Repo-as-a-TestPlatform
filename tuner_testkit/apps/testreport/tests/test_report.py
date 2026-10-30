from __future__ import annotations

import re
from datetime import date
from pathlib import Path

import pytest

from tuner_testkit.apps.testreport import cli
from tuner_testkit.apps.testreport.draft import TODO_MARK, compose
from tuner_testkit.apps.testreport.profile import Profile, load_profile
from tuner_testkit.apps.testreport.render import ReportFormatError, parse_report
from tuner_testkit.apps.testreport.stats import build

from .test_stats import NEW_COUNTS, NEW_DEFECTS, NEW_ROWS

PROFILE = """---
kind: domain-note
domain: testreport
---

# 演示系统测试报告静态信息

## 基本信息

| 项 | 内容 |
| --- | --- |
| 系统名称 | 演示系统 |
| 编制单位 | 演示公司 |
| 需方 | 甲方 |
| 开发方 | 乙方 |
| 测试类型 | 功能测试 |

## 系统简介

演示系统用于登录。

## 运行现场

浏览器访问。

## 引用文档

| 文档编号 | 文档标题 | 编写单位 | 发布日期 |
| --- | --- | --- | --- |
| DOC-1 | 演示系统需求规格说明 | 乙方 | 2026-10-01 |

## 软件环境

| 支持软件名称 | 版本号 | 用途 | 介质 | 来源 | 保密性 |
| --- | --- | --- | --- | --- | --- |
| Chrome | 120 | 浏览器 | 网络 | 官网 | 非密 |

## 硬件环境

| 硬件名称 | 配置 | 推荐型号 | 用途 | 数量 | 保密性 |
| --- | --- | --- | --- | --- | --- |
| PC | 16G | / | 测试机 | 1 | 非密 |

## 测试环境影响

| 测试环境 | 实装环境 | 影响分析 | 解决措施 |
| --- | --- | --- | --- |
| 符合标准的操作流程 | 可能出现无法预料的操作 | 场景覆盖不全 | 尽量覆盖可测场景 |

## 通过准则

| 项 | 内容 |
| --- | --- |
| 不允许遗留的缺陷等级 | 1级、2级 |
| 允许未执行用例 | 是 |
"""

TEMPLATE = Path(__file__).resolve().parents[1] / "templates" / "gjb438b.docx"


def _draft(sheet: Path, root: Path, profile: Profile) -> tuple[str, list[str]]:
    data = build([sheet], root=root)
    return compose(data, profile, slug="demo", sut="demo", env="dev", task_id="TASK-1", profile_rel="p.md", today=date(2026, 10, 29), root=root)


def _profile(tmp_path: Path) -> Profile:
    path = tmp_path / "profile.md"
    path.write_text(PROFILE, encoding="utf-8")
    return load_profile(path)


def _resolve_todos(text: str) -> str:
    return re.sub(r"<!-- TODO [^>]*-->", "无。", text)


def test_profile_parses_sections(tmp_path: Path) -> None:
    profile = _profile(tmp_path)
    assert profile.get("系统名称") == "演示系统"
    assert profile.intro == ["演示系统用于登录。"]
    assert profile.references[0]["文档编号"] == "DOC-1"
    assert profile.software[0]["支持软件名称"] == "Chrome"
    assert profile.criteria["允许未执行用例"] == "是"
    assert "软件标识号" in profile.missing()


def test_draft_without_profile_leaves_todos(sheet_factory, tmp_path: Path) -> None:
    sheet = sheet_factory(rows=NEW_ROWS, counts=NEW_COUNTS, total=5, defects=NEW_DEFECTS)
    text, hints = _draft(sheet, tmp_path, Profile())
    assert any("系统名称" in hint for hint in hints)
    assert any("编制单位" in hint for hint in hints)
    assert any("待评审" in hint for hint in hints)
    report = tmp_path / "report.md"
    report.write_text(text, encoding="utf-8")
    with pytest.raises(ReportFormatError, match="TODO"):
        parse_report(report)


def test_draft_numbers_and_verdict(sheet_factory, tmp_path: Path) -> None:
    sheet = sheet_factory(rows=NEW_ROWS, counts=NEW_COUNTS, total=5, defects=NEW_DEFECTS)
    text, hints = _draft(sheet, tmp_path, _profile(tmp_path))
    assert hints == ["按遗留缺陷与未执行用例写改进建议，没有建议写「无。」"]
    report = tmp_path / "report.md"
    report.write_text(_resolve_todos(text), encoding="utf-8")
    parsed = parse_report(report)
    assert parsed["system_name"] == "演示系统"
    assert parsed["verdict"] == "未通过测试"
    assert parsed["version"] == "1.2.3"
    assert parsed["s"]["period"] == "2026年10月28日～2026年10月29日"
    assert parsed["s"]["testers"] == "测试员"
    assert parsed["fr"] == {"designed": "5", "executed": "4", "passed": "1", "failed": "3", "not_run": "1"}
    assert parsed["rg"]["reuse"] == "2/0/0" and parsed["rg"]["executed"] == "1"
    assert [row["level"] for row in parsed["defect_rows"]] == ["1级", "2级", "3级", "4级"]
    assert parsed["dt"]["closed"] == "2" and parsed["dt"]["open"] == "2"
    assert {row["case"] for row in parsed["residuals"]} == {"Login-Open-locked", "Login-Wait-deleted"}
    assert parsed["references"][0]["code"] == "DOC-1"
    assert len(parsed["cases"]) == 5
    first = parsed["cases"][0]
    assert first["action"] == "1、输入账号和口令，点击登录\n2、刷新页面"
    assert first["pre"].startswith("账号已存在。\n测试数据：")
    assert {case["result"] for case in parsed["cases"]} == {"通过", "未通过", "未执行"}
    assert any(note.startswith("未执行或阻塞用例1个") for note in parsed["notes"])


def test_draft_reads_dotted_dates(sheet_factory, tmp_path: Path) -> None:
    sheet = sheet_factory(rows=NEW_ROWS, counts=NEW_COUNTS, total=5, defects=NEW_DEFECTS)
    path = sheet / "testsheet.md"
    path.write_text(path.read_text(encoding="utf-8").replace("2026-10-28 ~ 2026-10-29", "2026.10.21"), encoding="utf-8")
    text, hints = _draft(sheet, tmp_path, _profile(tmp_path))
    assert not any("日期" in hint for hint in hints)
    assert "| 测试周期 | 2026年10月21日 |" in text


def test_draft_basis_path_keeps_whole_and_takes_title(sheet_factory, tmp_path: Path) -> None:
    doc = tmp_path / "assets" / "testreport" / "260101示例" / "需求.md"
    doc.parent.mkdir(parents=True)
    doc.write_text("# 演示系统 — 需求说明\n\n| 项目 | 内容 |\n| --- | --- |\n| 建立日期 | 2026.10.21 |\n\n正文。\n", encoding="utf-8")
    sheet = sheet_factory(rows=NEW_ROWS, counts=NEW_COUNTS, total=5, defects=NEW_DEFECTS)
    path = sheet / "testsheet.md"
    basis = "依据：`assets/testreport/260101示例/需求.md`、用户说明 2026-10-28。该报告结论为通过。"
    path.write_text(path.read_text(encoding="utf-8").replace("依据：用户说明 2026-10-28。", basis), encoding="utf-8")
    text, _ = _draft(sheet, tmp_path, _profile(tmp_path))
    refs = text.split("## 2 引用文档", 1)[1].split("## 3", 1)[0]
    assert "| 3 | / | 演示系统 — 需求说明（assets/testreport/260101示例/需求.md） | / | 2026.10.21 |" in refs
    assert "| 4 | / | 用户说明 | / | 2026-10-28 |" in refs
    assert "| 5 |" not in refs
    assert "测试单的依据是演示系统 — 需求说明（assets/testreport/260101示例/需求.md）、用户说明 2026-10-28。该报告结论为通过。" in text


def test_profile_slash_row_means_left_blank(sheet_factory, tmp_path: Path) -> None:
    blank = PROFILE.replace("| Chrome | 120 | 浏览器 | 网络 | 官网 | 非密 |", "| / | / | / | / | / | / |")
    path = tmp_path / "profile.md"
    path.write_text(blank, encoding="utf-8")
    sheet = sheet_factory(rows=NEW_ROWS, counts=NEW_COUNTS, total=5, defects=NEW_DEFECTS)
    text, hints = _draft(sheet, tmp_path, load_profile(path))
    assert not any("软件环境" in hint for hint in hints)
    assert "| 1 | / | / | / | / | / | / |" in text.split("#### 3.2.1 软件环境", 1)[1].split("####", 1)[0]


def test_draft_refuses_sheet_with_errors(sheet_factory, tmp_path: Path) -> None:
    sheet = sheet_factory(rows=NEW_ROWS, counts={"通过": 9}, total=5, defects=NEW_DEFECTS)
    out = tmp_path / "out.md"
    code = cli.main(["draft", str(sheet), "--sut", "demo", "--slug", "demo", "--profile", str(tmp_path / "none.md"), "--out", str(out)])
    assert code == 1
    assert not out.exists()


def test_render_refuses_todos(sheet_factory, tmp_path: Path) -> None:
    sheet = sheet_factory(rows=NEW_ROWS, counts=NEW_COUNTS, total=5, defects=NEW_DEFECTS)
    text, _ = _draft(sheet, tmp_path, Profile())
    report = tmp_path / "report.md"
    report.write_text(text, encoding="utf-8")
    assert cli.main(["render", str(report)]) == 1


def test_cli_render_uses_packaged_template(sheet_factory, tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    pytest.importorskip("docxtpl")
    sheet = sheet_factory(rows=NEW_ROWS, counts=NEW_COUNTS, total=5, defects=NEW_DEFECTS)
    text, _ = _draft(sheet, tmp_path, _profile(tmp_path))
    report = tmp_path / "report.md"
    report.write_text(_resolve_todos(text), encoding="utf-8")
    assert cli.main(["render", str(report)]) == 0
    assert "check leftover_tags=0 case_rows=5 ok" in capsys.readouterr().out
    assert len(list((tmp_path / "artifacts" / "runs").glob("*testreport-demo/*.docx"))) == 1


def test_render_docx(sheet_factory, tmp_path: Path) -> None:
    pytest.importorskip("docxtpl")
    from docx import Document

    from tuner_testkit.apps.testreport.render import render_docx

    sheet = sheet_factory(rows=NEW_ROWS, counts=NEW_COUNTS, total=5, defects=NEW_DEFECTS)
    text, _ = _draft(sheet, tmp_path, _profile(tmp_path))
    report = tmp_path / "report.md"
    report.write_text(_resolve_todos(text), encoding="utf-8")
    out = tmp_path / "report.docx"
    summary = render_docx(report, TEMPLATE, out)
    assert summary["cases"] == 5 and summary["verdict"] == "未通过测试"
    assert summary["leftover_tags"] == 0 and summary["case_rows"] == 5

    doc = Document(str(out))
    body = "\n".join(p.text for p in doc.paragraphs)
    assert "{{" not in body and "{%" not in body
    assert "演示系统" in body and "a)需方：甲方" in body
    cases = next(t for t in doc.tables if "用例编号（测试标识）" in [c.text for c in t.rows[0].cells])
    assert len(cases.rows) == 1 + 5
    assert cases.rows[1].cells[3].text == "Login-Fixed-frozen"
    defects = next(t for t in doc.tables if t.rows[0].cells[0].text == "严重性等级")
    assert [row.cells[0].text for row in defects.rows[2:]] == ["1级", "2级", "3级", "4级", "合计"]
    for table in doc.tables:
        for row in table.rows:
            for c in row.cells:
                assert TODO_MARK not in c.text and "{{" not in c.text
