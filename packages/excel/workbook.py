"""Thin, business-agnostic wrapper around openpyxl for test asset reading.

Design contract (see :mod:`packages.excel` docstring for the full scope):

- accept bytes or a local path and return stringified cell values;
- preserve row order (including the header row);
- strip leading/trailing whitespace from header cell values only (data cells
  are kept verbatim so tests can spot rogue whitespace in payloads);
- fall back to CSV parsing when the bytes are not a valid xlsx container
  (magic bytes ``PK\\x03\\x04`` / ``PK\\x05\\x06`` / ``PK\\x07\\x08``).

This file contains zero business logic — no 授权/产线/省份 knowledge. The
caller is responsible for mapping the stringified cells onto its domain
model (e.g. :class:`packages.argon.permission.models.AuthorizationRecord`).
"""
from __future__ import annotations

import csv
import io
from dataclasses import dataclass
from pathlib import Path
from typing import Iterator

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException


_XLSX_MAGICS: tuple[bytes, ...] = (b"PK\x03\x04", b"PK\x05\x06", b"PK\x07\x08")


class ExcelError(RuntimeError):
    """Raised when the workbook cannot be parsed as xlsx or CSV."""


def _stringify(value: object) -> str:
    """Return the cell value as ``str`` (``""`` when the cell is empty).

    No type coercion beyond ``str(value)`` — the goal is to keep raw content
    available for the caller without committing to a date / number format.
    """
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    return str(value)


@dataclass(frozen=True)
class ExcelSheet:
    """Immutable view on a single sheet (name + rows).

    ``rows`` is the raw grid: the first row is typically the header, the
    rest are data rows. Use :meth:`as_records` when the first row holds
    the header names.
    """

    name: str
    rows: tuple[tuple[str, ...], ...]

    def __len__(self) -> int:
        return len(self.rows)

    def __iter__(self) -> Iterator[tuple[str, ...]]:
        return iter(self.rows)

    @property
    def header(self) -> tuple[str, ...]:
        if not self.rows:
            return ()
        return tuple(cell.strip() for cell in self.rows[0])

    @property
    def data_rows(self) -> tuple[tuple[str, ...], ...]:
        if not self.rows:
            return ()
        return self.rows[1:]

    def as_records(self) -> list[dict[str, str]]:
        """Return ``[{header: value, ...}, ...]`` for the non-header rows.

        Missing cells are filled with ``""`` so every record has the same
        keys. Extra cells beyond the header length are dropped silently —
        this matches the way openpyxl pads rows to the sheet dimension.
        """
        hdr = self.header
        if not hdr:
            return []
        out: list[dict[str, str]] = []
        for row in self.data_rows:
            record = {}
            for idx, key in enumerate(hdr):
                record[key] = row[idx] if idx < len(row) else ""
            out.append(record)
        return out


@dataclass(frozen=True)
class ExcelWorkbook:
    """Immutable view on a parsed workbook (sheet-name order preserved)."""

    sheets: tuple[ExcelSheet, ...]

    @classmethod
    def from_bytes(cls, data: bytes | bytearray) -> "ExcelWorkbook":
        """Parse an in-memory byte blob. Falls back to CSV when not xlsx."""
        if data is None:
            raise ExcelError("empty data")
        buf = bytes(data)
        if not buf:
            raise ExcelError("empty data")
        if _looks_like_xlsx(buf):
            return _load_xlsx_bytes(buf)
        return _load_csv_bytes(buf)

    @classmethod
    def from_path(cls, path: str | Path) -> "ExcelWorkbook":
        """Parse a local file; dispatches by extension first, then by magic."""
        p = Path(path)
        if not p.exists():
            raise ExcelError(f"file not found: {p}")
        if p.suffix.lower() == ".csv":
            return _load_csv_bytes(p.read_bytes())
        return cls.from_bytes(p.read_bytes())

    @property
    def first_sheet(self) -> ExcelSheet:
        if not self.sheets:
            raise ExcelError("workbook has no sheets")
        return self.sheets[0]

    def sheet(self, name: str) -> ExcelSheet:
        for s in self.sheets:
            if s.name == name:
                return s
        raise ExcelError(f"sheet {name!r} not found; available: {[s.name for s in self.sheets]}")

    def first_sheet_rows(self) -> list[dict[str, str]]:
        """Convenience: return the first sheet's data rows as ``list[dict]``.

        Header names are trimmed (leading/trailing whitespace removed) so
        callers can look up columns by the logical header text without
        fighting spreadsheet artefacts.
        """
        return self.first_sheet.as_records()


# ---------------------------------------------------------------------------
# Top-level helpers (thin wrappers — no business logic)
# ---------------------------------------------------------------------------
def read_rows(source: bytes | str | Path, *, sheet: str | None = None) -> list[list[str]]:
    """Return stringified rows (header + data) for one sheet of the workbook.

    ``source`` may be bytes (in-memory blob) or a path-like. When ``sheet``
    is omitted, the first sheet is returned.
    """
    wb = _load_any(source)
    target = wb.sheet(sheet) if sheet else wb.first_sheet
    return [list(row) for row in target.rows]


def first_sheet_rows(source: bytes | str | Path) -> list[dict[str, str]]:
    """Return the first sheet as ``list[{header: value}]``."""
    return _load_any(source).first_sheet_rows()


def rows_as_records(
    source: bytes | str | Path, *, sheet: str | None = None
) -> list[dict[str, str]]:
    """Return any named sheet as ``list[{header: value}]`` records."""
    wb = _load_any(source)
    target = wb.sheet(sheet) if sheet else wb.first_sheet
    return target.as_records()


# ---------------------------------------------------------------------------
# Internal loaders
# ---------------------------------------------------------------------------
def _load_any(source: bytes | str | Path) -> ExcelWorkbook:
    if isinstance(source, (bytes, bytearray)):
        return ExcelWorkbook.from_bytes(source)
    return ExcelWorkbook.from_path(source)


def _looks_like_xlsx(buf: bytes) -> bool:
    return any(buf.startswith(m) for m in _XLSX_MAGICS)


def _load_xlsx_bytes(buf: bytes) -> ExcelWorkbook:
    try:
        wb = load_workbook(io.BytesIO(buf), read_only=True, data_only=True)
    except InvalidFileException as e:
        raise ExcelError(f"not a valid xlsx: {e}") from e

    sheets: list[ExcelSheet] = []
    for name in wb.sheetnames:
        ws = wb[name]
        rows: list[tuple[str, ...]] = []
        for raw_row in ws.iter_rows(values_only=True):
            rows.append(tuple(_stringify(cell) for cell in raw_row))
        sheets.append(ExcelSheet(name=name, rows=tuple(rows)))
    wb.close()
    return ExcelWorkbook(sheets=tuple(sheets))


def _load_csv_bytes(buf: bytes) -> ExcelWorkbook:
    try:
        text = buf.decode("utf-8-sig")
    except UnicodeDecodeError:
        try:
            text = buf.decode("gbk")
        except UnicodeDecodeError as e:
            raise ExcelError(f"csv decode failed (tried utf-8 and gbk): {e}") from e
    reader = csv.reader(io.StringIO(text))
    rows = tuple(tuple(_stringify(c) for c in row) for row in reader)
    return ExcelWorkbook(sheets=(ExcelSheet(name="csv", rows=rows),))


__all__ = [
    "ExcelError",
    "ExcelSheet",
    "ExcelWorkbook",
    "first_sheet_rows",
    "read_rows",
    "rows_as_records",
]
